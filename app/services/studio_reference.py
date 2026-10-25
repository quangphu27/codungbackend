"""Ảnh trang sách + PDF/DOCX — ngữ cảnh cho Teacher Studio AI."""
from __future__ import annotations

import base64
import io
import re

from flask import current_app
from werkzeug.datastructures import FileStorage

MAX_IMAGES = 8
MAX_TEXT_CHARS = 14_000
MAX_FILE_BYTES = 12 * 1024 * 1024
ALLOWED_IMAGE = {"image/jpeg", "image/jpg", "image/png", "image/webp"}
ALLOWED_DOC = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}


def _read_upload_bytes(upload: FileStorage) -> bytes:
    upload.stream.seek(0)
    data = upload.read()
    upload.stream.seek(0)
    return data or b""


def _guess_mime(upload: FileStorage, raw: bytes) -> str:
    mime = (upload.mimetype or "").split(";")[0].strip().lower()
    name = (upload.filename or "").lower()
    if mime in ALLOWED_IMAGE or mime in ALLOWED_DOC:
        return mime
    if name.endswith((".jpg", ".jpeg")):
        return "image/jpeg"
    if name.endswith(".png"):
        return "image/png"
    if name.endswith(".webp"):
        return "image/webp"
    if name.endswith(".pdf"):
        return "application/pdf"
    if name.endswith(".docx"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if name.endswith(".doc"):
        return "application/msword"
    if raw[:4] == b"%PDF":
        return "application/pdf"
    if raw[:2] == b"PK":
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return mime or "application/octet-stream"


def _resize_image_b64(raw: bytes, mime: str) -> tuple[str, str]:
    from PIL import Image

    img = Image.open(io.BytesIO(raw))
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    w, h = img.size
    max_side = 1400
    if max(w, h) > max_side:
        scale = max_side / float(max(w, h))
        img = img.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85, optimize=True)
    return "image/jpeg", base64.b64encode(buf.getvalue()).decode("ascii")


def _extract_docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf_text(raw: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    reader = PdfReader(io.BytesIO(raw))
    chunks = []
    for page in reader.pages[:40]:
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        t = re.sub(r"\s+", " ", t).strip()
        if t:
            chunks.append(t)
    return "\n\n".join(chunks)


def process_reference_uploads(uploads: list[FileStorage]) -> dict:
    """Trả về book_reference { text, images: [{mime, data}], filenames }."""
    text_parts: list[str] = []
    images: list[dict] = []
    filenames: list[str] = []

    for upload in uploads:
        if not upload or not upload.filename:
            continue
        raw = _read_upload_bytes(upload)
        if not raw or len(raw) > MAX_FILE_BYTES:
            raise ValueError(f"File quá lớn (tối đa {MAX_FILE_BYTES // (1024 * 1024)} MB): {upload.filename}")
        mime = _guess_mime(upload, raw)
        filenames.append(upload.filename)

        if mime in ALLOWED_IMAGE or mime.startswith("image/"):
            if len(images) >= MAX_IMAGES:
                continue
            try:
                out_mime, b64 = _resize_image_b64(raw, mime)
                images.append({"mime": out_mime, "data": b64})
            except Exception:
                current_app.logger.exception("studio reference image %s", upload.filename)
                raise ValueError(f"Không đọc được ảnh: {upload.filename}") from None
        elif mime == "application/pdf" or (upload.filename or "").lower().endswith(".pdf"):
            extracted = _extract_pdf_text(raw)
            if extracted:
                text_parts.append(f"[PDF: {upload.filename}]\n{extracted}")
            elif len(images) < MAX_IMAGES:
                raise ValueError(
                    f"PDF «{upload.filename}» không trích được chữ — hãy chụp ảnh trang sách hoặc cài pypdf trên server."
                )
        elif "wordprocessingml" in mime or (upload.filename or "").lower().endswith((".docx", ".doc")):
            if (upload.filename or "").lower().endswith(".doc"):
                raise ValueError("File .doc cũ — hãy lưu sang .docx hoặc PDF.")
            extracted = _extract_docx_text(raw)
            if extracted:
                text_parts.append(f"[Word: {upload.filename}]\n{extracted}")
            else:
                raise ValueError(f"File Word trống hoặc không đọc được: {upload.filename}")
        else:
            raise ValueError(f"Không hỗ trợ loại file: {upload.filename}")

    full_text = "\n\n".join(text_parts).strip()
    if len(full_text) > MAX_TEXT_CHARS:
        full_text = full_text[:MAX_TEXT_CHARS] + "\n… (đã rút gọn)"

    if not full_text and not images:
        raise ValueError("Không có nội dung — chọn ảnh trang sách hoặc file PDF/Word.")

    return {
        "text": full_text,
        "images": images,
        "filenames": filenames,
        "image_count": len(images),
        "has_text": bool(full_text),
    }


def book_reference_prompt_block(book_reference: dict | None) -> tuple[str, list[dict]]:
    """Đoạn prompt + parts Gemini (inline_data)."""
    if not book_reference or not isinstance(book_reference, dict):
        return "", []
    text = (book_reference.get("text") or "").strip()
    images = book_reference.get("images") if isinstance(book_reference.get("images"), list) else []
    media_parts = []
    for img in images[:MAX_IMAGES]:
        if not isinstance(img, dict):
            continue
        data = img.get("data")
        mime = img.get("mime") or "image/jpeg"
        if data:
            media_parts.append({"inline_data": {"mime_type": mime, "data": data}})
    block = ""
    if text:
        block = (
            "\n\n=== TÀI LIỆU / SÁCH GIÁO VIÊN ĐÍNH KÈM (BẮT BUỘC bám sát nội dung này) ===\n"
            f"{text}\n=== HẾT TÀI LIỆU ===\n"
        )
    if media_parts:
        block += (
            f"\nCó {len(media_parts)} ảnh trang sách — đọc chữ trong ảnh, nhận diện đúng bài "
            "và soạn giáo án/đề/slide khớp sách.\n"
        )
    return block, media_parts


def sanitize_book_reference_for_storage(book_reference: dict | None) -> dict | None:
    """Bỏ base64 khi lưu draft (chỉ metadata)."""
    if not book_reference or not isinstance(book_reference, dict):
        return None
    return {
        "text": (book_reference.get("text") or "")[:2000],
        "filenames": book_reference.get("filenames") or [],
        "image_count": book_reference.get("image_count") or len(book_reference.get("images") or []),
        "has_text": bool(book_reference.get("text")),
    }
