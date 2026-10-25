"""
Professional lesson-plan DOCX builder for Teacher Studio.
Clear structure, bold headings, font sizes, tables, alignment.
"""
from __future__ import annotations

import io
from typing import Any


NAVY = (0x0F, 0x27, 0x44)
ORANGE = (0xF3, 0x9C, 0x4A)
TEAL = (0x1A, 0x9B, 0x8E)
GRAY = (0x4B, 0x55, 0x63)


def _parse_hex(value, fallback=NAVY):
    text = str(value or "").strip().lstrip("#")
    if len(text) == 3:
        text = "".join(ch * 2 for ch in text)
    if len(text) != 6:
        return fallback
    try:
        return tuple(int(text[i : i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return fallback


def _hex_of(rgb) -> str:
    return "{:02X}{:02X}{:02X}".format(*rgb)


def _theme_from(doc_or_dict: dict | None):
    data = doc_or_dict or {}
    navy = _parse_hex(data.get("theme_color") or data.get("primary_color"), NAVY)
    orange = _parse_hex(data.get("theme_accent") or data.get("accent_color"), ORANGE)
    teal = (
        min(255, int(navy[0] * 0.25 + 26)),
        min(255, int(navy[1] * 0.35 + 120)),
        min(255, int(navy[2] * 0.4 + 110)),
    )
    return {"navy": navy, "orange": orange, "teal": teal, "gray": GRAY}


def _rgb(tup):
    from docx.shared import RGBColor
    return RGBColor(*tup)


def _set_run(run, *, size=12, bold=False, italic=False, color=None, name="Times New Roman"):
    from docx.shared import Pt
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = _rgb(color)
    # Ensure East Asian font compatibility on Windows/WPS
    try:
        r = run._element
        rPr = r.get_or_add_rPr()
        from docx.oxml.ns import qn
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), name)
        rFonts.set(qn("w:hAnsi"), name)
        rFonts.set(qn("w:eastAsia"), "Times New Roman")
    except Exception:
        pass


def _add_para(doc, text="", *, size=12, bold=False, italic=False, color=None,
              align="left", space_before=0, space_after=6, first_line=False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Pt, Twips

    p = doc.add_paragraph()
    align_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    p.alignment = align_map.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_line:
        pf.first_line_indent = Pt(18)
    run = p.add_run(str(text or ""))
    _set_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def _add_heading_custom(doc, text, level=1, theme=None):
    """Styled section headings (not relying on Word built-in styles alone)."""
    theme = theme or {"navy": NAVY, "teal": TEAL, "gray": GRAY}
    sizes = {0: 18, 1: 14, 2: 13, 3: 12}
    colors = {0: theme["navy"], 1: theme["navy"], 2: theme.get("teal", TEAL), 3: theme.get("gray", GRAY)}
    align = "center" if level == 0 else "left"
    before = 14 if level <= 1 else 10
    p = _add_para(
        doc, text,
        size=sizes.get(level, 12),
        bold=True,
        color=colors.get(level, theme["navy"]),
        align=align,
        space_before=before,
        space_after=6,
    )
    if level == 1:
        try:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "12")
            bottom.set(qn("w:space"), "4")
            bottom.set(qn("w:color"), _hex_of(theme["navy"]))
            pBdr.append(bottom)
            pPr.append(pBdr)
        except Exception:
            pass
    return p


def _add_bullet(doc, text, *, size=12, bold_prefix=None):
    from docx.shared import Pt
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    if bold_prefix:
        r1 = p.add_run(str(bold_prefix))
        _set_run(r1, size=size, bold=True, color=NAVY)
        r2 = p.add_run(str(text))
        _set_run(r2, size=size, color=GRAY)
    else:
        r = p.add_run(str(text))
        _set_run(r, size=size, color=GRAY)
    return p


def _set_cell_shading(cell, hex_color: str):
    try:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), hex_color)
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
    except Exception:
        pass


def _fill_cell(cell, text, *, bold=False, size=11, color=None, center=False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text or ""))
    _set_run(run, size=size, bold=bold, color=color or NAVY)


def _add_info_table(doc, rows: list[tuple[str, str]], theme=None):
    """2-column info table (Label | Value)."""
    from docx.shared import Cm

    theme = theme or {"navy": NAVY}
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.autofit = True
    for i, (label, value) in enumerate(rows):
        _fill_cell(table.cell(i, 0), label, bold=True, size=11, color=theme["navy"])
        _fill_cell(table.cell(i, 1), value, size=11, color=GRAY)
        _set_cell_shading(table.cell(i, 0), "E8F1FB")
        table.cell(i, 0).width = Cm(4.5)
        table.cell(i, 1).width = Cm(12)
    doc.add_paragraph()


def _add_data_table(doc, headers: list[str], rows: list[list[Any]], theme=None):
    theme = theme or {"navy": NAVY}
    cols = max(len(headers), 1)
    table = doc.add_table(rows=1 + len(rows), cols=cols)
    table.style = "Table Grid"
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        _fill_cell(cell, h, bold=True, size=11, color=(0xFF, 0xFF, 0xFF), center=True)
        _set_cell_shading(cell, _hex_of(theme["navy"]))
    for r_idx, row in enumerate(rows, start=1):
        for c in range(cols):
            val = row[c] if c < len(row) else ""
            cell = table.cell(r_idx, c)
            _fill_cell(cell, val, size=11, color=GRAY)
            if r_idx % 2 == 0:
                _set_cell_shading(cell, "F7FAFC")
    doc.add_paragraph()


def _procedure_rows(lesson: dict) -> list[list[str]]:
    """Build teaching procedure table from sections."""
    rows = []
    default_times = ["5'", "12'", "15'", "10'", "8'", "5'"]
    for i, sec in enumerate(lesson.get("sections") or []):
        heading = sec.get("heading_vi") or sec.get("heading_en") or f"Step {i+1}"
        content = (sec.get("content_vi") or sec.get("content_en") or "").strip()
        # shorten for table cell
        if len(content) > 360:
            content = content[:357] + "..."
        time = sec.get("duration") or (default_times[i] if i < len(default_times) else "8'")
        teacher = sec.get("teacher_activity") or "Hướng dẫn / Facilitate"
        student = sec.get("student_activity") or "Thực hiện / Practice"
        rows.append([str(time), str(heading), str(content), str(teacher), str(student)])
    return rows


def _is_clil_lesson(lesson: dict) -> bool:
    return (
        (lesson or {}).get("lesson_format") == "edubridge_clil"
        or isinstance((lesson or {}).get("activities"), list)
    )


def _build_edubridge_clil_docx_bytes(lesson: dict) -> bytes:
    """Giáo án tiếng Anh (CLIL) + phụ lục Teacher's English Support Notes."""
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    theme = _theme_from(lesson)
    header = lesson.get("lesson_header") if isinstance(lesson.get("lesson_header"), dict) else {}
    title = lesson.get("title_en") or lesson.get("title_vi") or "Lesson plan"
    grade = lesson.get("grade") or ""
    subject = lesson.get("subject") or ""
    topic = lesson.get("topic") or ""
    duration = lesson.get("duration_minutes") or 45

    _add_para(doc, "EDUBRIDGE · CLIL LESSON PLAN", size=11, bold=True, color=theme["orange"], align="center", space_after=2)
    _add_para(doc, "LESSON PLAN", size=14, bold=True, color=theme["navy"], align="center", space_after=6)
    _add_heading_custom(doc, title, level=0, theme=theme)

    info_rows = [
        ("Subject", subject),
        ("Grade", grade),
        ("Unit / Topic", header.get("unit_topic") or topic),
        ("Lesson", header.get("lesson_title") or title),
        ("Week", header.get("week") or "—"),
        ("Period", header.get("period") or "—"),
        ("Duration", f"{duration} minutes"),
        ("Textbook / Source", header.get("textbook_source") or "Textbook / teacher materials"),
        ("Teacher English level", header.get("teacher_english_level") or lesson.get("teacher_english_level") or "A2"),
        ("Students' English level", header.get("student_english_level") or lesson.get("student_english_level") or "A2"),
        ("Class size", header.get("class_size") or "—"),
        ("Equipment", header.get("equipment") or "Projector, board"),
        ("Approach", header.get("teaching_approach") or "CLIL, active learning"),
    ]
    _add_heading_custom(doc, "LESSON INFORMATION", level=1, theme=theme)
    _add_info_table(doc, info_rows, theme=theme)

    desc = (lesson.get("description") or "").strip()
    if desc:
        _add_para(doc, "Description", size=12, bold=True, color=theme["teal"], space_after=4)
        _add_para(doc, desc, size=12, color=GRAY, align="justify", first_line=True, space_after=8)

    src_notes = (lesson.get("source_notes") or "").strip()
    if src_notes:
        _add_para(doc, "Source / material notes", size=11, bold=True, color=theme["orange"], space_after=4)
        _add_para(doc, src_notes, size=11, italic=True, color=GRAY, space_after=8)

    obj = lesson.get("objectives") if isinstance(lesson.get("objectives"), dict) else {}
    _add_heading_custom(doc, "I. OBJECTIVES", level=1, theme=theme)
    _add_para(doc, "By the end of this lesson, students will be able to:", size=12, bold=True, space_after=4)
    for label, key in (
        ("1. Knowledge", "knowledge"),
        ("2. Competences", "competences"),
        ("3. Personal qualities", "personal_qualities"),
        ("4. Digital competence", "digital_competence"),
    ):
        items = obj.get(key)
        if key == "digital_competence" and not items:
            continue
        if isinstance(items, str):
            items = [items]
        if not items:
            continue
        _add_para(doc, label, size=11, bold=True, color=theme["navy"], space_before=6, space_after=2)
        for it in items:
            _add_bullet(doc, f" {it}")

    _add_heading_custom(doc, "II. MATERIALS", level=1, theme=theme)
    mats = lesson.get("materials")
    if isinstance(mats, dict):
        for label, key in (
            ("Textbook / Learning materials", "textbook"),
            ("Equipment", "equipment"),
            ("Worksheets", "worksheets"),
            ("Digital resources", "digital"),
            ("Other", "other"),
        ):
            items = mats.get(key) or []
            if items:
                _add_para(doc, label, size=11, bold=True, color=theme["teal"], space_after=2)
                for m in items:
                    _add_bullet(doc, f" {m}")
    elif isinstance(mats, list):
        for m in mats:
            _add_bullet(doc, f" {m}")

    analysis = (lesson.get("subject_language_analysis") or lesson.get("core_content_en") or "").strip()
    if analysis:
        _add_para(doc, "Subject content / Language analysis", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
        _add_para(doc, analysis, size=12, color=GRAY, align="justify", first_line=True, space_after=8)

    vocab = lesson.get("vocabulary") or []
    if vocab:
        rows = []
        for v in vocab:
            if not isinstance(v, dict):
                continue
            rows.append([
                v.get("term") or "",
                v.get("ipa") or "",
                v.get("meaning_vi") or v.get("meaning") or "",
                v.get("explanation_en") or "",
                v.get("example") or "",
            ])
        _add_para(doc, "Key vocabulary", size=12, bold=True, color=theme["teal"], space_before=6, space_after=4)
        _add_data_table(
            doc,
            ["English term", "IPA", "Vietnamese", "Simple English explanation", "Example"],
            rows,
            theme=theme,
        )

    diffs = lesson.get("anticipated_difficulties") or []
    if diffs:
        rows = []
        for d in diffs:
            if isinstance(d, dict):
                rows.append([d.get("difficulty") or "", d.get("solution") or ""])
        if rows:
            _add_para(doc, "Assumptions / Anticipated difficulties", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
            _add_data_table(doc, ["Anticipated difficulties", "Suggested solutions"], rows, theme=theme)

    _add_heading_custom(doc, "III. PROCEDURES", level=1, theme=theme)
    activities = lesson.get("activities") or []
    if not activities:
        activities = [
            {**s, "name": s.get("heading_en") or s.get("heading_vi")}
            for s in (lesson.get("sections") or [])
            if isinstance(s, dict)
        ]
    for idx, act in enumerate(activities, 1):
        if not isinstance(act, dict):
            continue
        name = act.get("name") or f"Activity {idx}"
        _add_heading_custom(doc, f"{idx}. {name}", level=2, theme=theme)
        if act.get("duration"):
            _add_para(doc, f"Duration: {act['duration']}", size=11, italic=True, color=theme["orange"], space_after=4)
        for label, key in (
            ("a. Objectives", "objectives"),
            ("b. Content", "content"),
            ("c. Expected outcomes", "expected_outcomes"),
        ):
            val = (act.get(key) or "").strip()
            if val:
                _add_para(doc, label, size=11, bold=True, color=theme["navy"], space_after=2)
                _add_para(doc, val, size=12, color=GRAY, align="justify", first_line=True, space_after=4)
        org = act.get("organisation") or []
        if org:
            _add_para(doc, "d. Organisation", size=11, bold=True, color=theme["navy"], space_after=4)
            rows = []
            for row in org:
                if isinstance(row, dict):
                    rows.append([
                        row.get("teacher_en") or "",
                        row.get("student_en") or "",
                        row.get("contents") or "",
                    ])
            if rows:
                _add_data_table(
                    doc,
                    ["TEACHER'S ACTIVITIES", "STUDENTS' ACTIVITIES", "CONTENTS"],
                    rows,
                    theme=theme,
                )
        if act.get("assessment") or act.get("activity_assessment"):
            _add_para(doc, "e. Assessment", size=11, bold=True, color=theme["navy"], space_after=2)
            _add_para(doc, act.get("assessment") or act.get("activity_assessment"), size=12, color=GRAY, space_after=6)

    exercises = lesson.get("exercises") or []
    if exercises:
        _add_heading_custom(doc, "IV. EXERCISES & ASSESSMENT", level=1, theme=theme)
        for ex in exercises:
            if not isinstance(ex, dict):
                continue
            if ex.get("instructions_en"):
                _add_para(doc, ex["instructions_en"], size=12, bold=True, color=theme["teal"], space_after=4)
            for t in ex.get("tasks") or []:
                _add_bullet(doc, f" {t}")
            if ex.get("answer_key"):
                _add_para(doc, f"Answer key: {ex['answer_key']}", size=11, italic=True, color=GRAY, space_after=4)

    notes = lesson.get("english_support_notes")
    if isinstance(notes, dict):
        doc.add_page_break()
        _add_heading_custom(doc, "TEACHER'S ENGLISH SUPPORT NOTES", level=1, theme=theme)
        _add_para(
            doc,
            "Appendix for teacher preparation — not part of the formal lesson plan table.",
            size=10, italic=True, color=GRAY, space_after=8,
        )
        expr = notes.get("essential_expressions") or []
        if expr:
            _add_para(doc, "A. Essential teaching expressions", size=12, bold=True, color=theme["teal"], space_after=4)
            rows = [[e.get("en") or "", e.get("vi") or ""] for e in expr if isinstance(e, dict)]
            if rows:
                _add_data_table(doc, ["English", "Vietnamese"], rows, theme=theme)
        terms = notes.get("difficult_terminology") or []
        if terms:
            _add_para(doc, "B. Difficult subject terminology", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
            rows = [[t.get("term") or "", t.get("ipa") or "", t.get("vi") or ""] for t in terms if isinstance(t, dict)]
            if rows:
                _add_data_table(doc, ["Term", "IPA", "Vietnamese"], rows, theme=theme)
        pron = notes.get("pronunciation_support") or []
        if pron:
            _add_para(doc, "C. Pronunciation support", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
            for p in pron:
                if isinstance(p, dict):
                    _add_bullet(doc, f" {p.get('phrase') or ''} {p.get('ipa') or ''} — {p.get('note') or ''}")
        simple = notes.get("simplified_explanations") or []
        if simple:
            _add_para(doc, "D. Simplified explanations", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
            for s in simple:
                if isinstance(s, dict):
                    _add_para(doc, s.get("concept") or "", size=11, bold=True, space_after=2)
                    if s.get("standard_en"):
                        _add_para(doc, f"Standard: {s['standard_en']}", size=11, color=GRAY, space_after=2)
                    if s.get("simple_en"):
                        _add_para(doc, f"Simple: {s['simple_en']}", size=11, color=GRAY, space_after=2)
                    if s.get("vietnamese_support"):
                        _add_para(doc, f"Vietnamese (for teacher): {s['vietnamese_support']}", size=11, italic=True, color=GRAY, space_after=4)
        sit = notes.get("classroom_situations") or []
        if sit:
            _add_para(doc, "E. Classroom situations", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
            rows = [[s.get("situation") or "", s.get("english_response") or ""] for s in sit if isinstance(s, dict)]
            if rows:
                _add_data_table(doc, ["Situation", "Suggested English response"], rows, theme=theme)

    _add_para(doc, "", space_after=8)
    _add_para(
        doc,
        "— EduBridge · Teacher Studio · CLIL lesson plan —",
        size=9, italic=True, color=theme["orange"], align="center", space_before=12,
    )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_rich_docx_bytes(lesson: dict) -> bytes:
    if _is_clil_lesson(lesson):
        return _build_edubridge_clil_docx_bytes(lesson)
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    theme = _theme_from(lesson)
    title_vi = lesson.get("title_vi") or lesson.get("title") or "Giáo án giảng dạy"
    title_en = lesson.get("title_en") or ""
    grade = lesson.get("grade") or ""
    subject = lesson.get("subject") or ""
    topic = lesson.get("topic") or ""
    description = (lesson.get("description") or "").strip()
    duration = lesson.get("duration_minutes") or lesson.get("duration") or 45
    materials = lesson.get("materials") or []
    if isinstance(materials, str):
        materials = [materials]

    # ===== HEADER =====
    _add_para(doc, "EDUBRIDGE · TEACHING AND LEARNING SUBJECTS IN ENGLISH", size=11, bold=True, color=theme["orange"], align="center", space_after=2)
    _add_para(doc, "GIÁO ÁN GIẢNG DẠY / BILINGUAL LESSON PLAN", size=13, bold=True, color=theme["navy"], align="center", space_after=4)
    _add_heading_custom(doc, title_vi, level=0, theme=theme)
    if title_en and title_en != title_vi:
        _add_para(doc, title_en, size=13, bold=True, italic=True, color=theme["teal"], align="center", space_after=10)

    # ===== INFO TABLE =====
    _add_heading_custom(doc, "I. THÔNG TIN CHUNG / GENERAL INFORMATION", level=1, theme=theme)
    _add_info_table(doc, [
        ("Khối / Grade", grade),
        ("Môn / Subject", subject),
        ("Chủ đề / Topic", topic),
        ("Thời lượng / Duration", f"{duration} phút"),
        ("Hình thức / Mode", lesson.get("mode") or "Lớp học · Giảng dạy theo chủ đề"),
        ("Ngôn ngữ / Language", "Song ngữ Việt – Anh / Bilingual VI–EN"),
    ], theme=theme)
    if description:
        _add_para(doc, "Mô tả / Description", size=12, bold=True, color=theme["teal"], space_after=4)
        _add_para(doc, description, size=12, color=GRAY, align="justify", first_line=True, space_after=8)

    # ===== OBJECTIVES =====
    _add_heading_custom(doc, "II. MỤC TIÊU BÀI HỌC / LEARNING OBJECTIVES", level=1, theme=theme)
    objs_vi = lesson.get("objectives_vi") or []
    objs_en = lesson.get("objectives_en") or []
    if objs_vi:
        _add_para(doc, "Mục tiêu (Tiếng Việt)", size=12, bold=True, color=theme["teal"], space_before=4, space_after=4)
        for i, o in enumerate(objs_vi, 1):
            _add_bullet(doc, f" {o}", bold_prefix=f"{i}.")
    if objs_en:
        _add_para(doc, "Objectives (English)", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
        for i, o in enumerate(objs_en, 1):
            _add_bullet(doc, f" {o}", bold_prefix=f"{i}.")
    knowledge = lesson.get("knowledge_vi") or []
    skills = lesson.get("skills_vi") or []
    qualities = lesson.get("qualities_vi") or []
    knowledge_en = lesson.get("knowledge_en") or []
    skills_en = lesson.get("skills_en") or []
    qualities_en = lesson.get("qualities_en") or []
    if knowledge or skills or qualities or knowledge_en or skills_en or qualities_en:
        _add_para(doc, "Kiến thức – Kỹ năng – Phẩm chất / Knowledge – Skills – Qualities", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
        if knowledge:
            _add_para(doc, "Kiến thức / Knowledge (VI)", size=11, bold=True, color=theme["navy"], space_after=2)
            for item in knowledge:
                _add_bullet(doc, f" {item}")
        if knowledge_en:
            _add_para(doc, "Knowledge (EN)", size=11, bold=True, color=theme["teal"], space_before=6, space_after=2)
            for item in knowledge_en:
                _add_bullet(doc, f" {item}")
        if skills:
            _add_para(doc, "Kỹ năng / Skills (VI)", size=11, bold=True, color=theme["navy"], space_before=6, space_after=2)
            for item in skills:
                _add_bullet(doc, f" {item}")
        if skills_en:
            _add_para(doc, "Skills (EN)", size=11, bold=True, color=theme["teal"], space_before=6, space_after=2)
            for item in skills_en:
                _add_bullet(doc, f" {item}")
        if qualities:
            _add_para(doc, "Phẩm chất / Qualities (VI)", size=11, bold=True, color=theme["navy"], space_before=6, space_after=2)
            for item in qualities:
                _add_bullet(doc, f" {item}")
        if qualities_en:
            _add_para(doc, "Qualities (EN)", size=11, bold=True, color=theme["teal"], space_before=6, space_after=2)
            for item in qualities_en:
                _add_bullet(doc, f" {item}")
    core_vi = (lesson.get("core_content_vi") or "").strip()
    core_en = (lesson.get("core_content_en") or "").strip()
    if core_vi or core_en:
        _add_para(doc, "Nội dung kiến thức cốt lõi", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
        if core_vi:
            _add_para(doc, core_vi, size=13, color=GRAY, align="justify", first_line=True, space_after=6)
        if core_en:
            _add_para(doc, core_en, size=12, italic=True, color=GRAY, align="justify", space_after=6)

    # ===== MATERIALS =====
    _add_heading_custom(doc, "III. ĐỒ DÙNG DẠY HỌC / MATERIALS", level=1, theme=theme)
    if materials:
        for m in materials:
            _add_bullet(doc, f" {m}")
    else:
        for m in [
            "Projector / slides",
            "Worksheet / handout",
            "Board & markers",
            "Group activity cards",
        ]:
            _add_bullet(doc, f" {m}")

    # ===== PROCEDURE TABLE =====
    _add_heading_custom(doc, "IV. TIẾN TRÌNH DẠY HỌC / TEACHING PROCEDURE", level=1, theme=theme)
    _add_para(
        doc,
        "Theo GDPT 2018: Mở đầu (Khởi động) → Hình thành kiến thức mới → Luyện tập → Vận dụng.",
        size=11, italic=True, color=GRAY, space_after=4,
    )
    _add_para(
        doc,
        "Bảng tiến trình (thời gian – hoạt động – nội dung – giáo viên – học sinh).",
        size=11, italic=True, color=GRAY, space_after=6,
    )
    proc = _procedure_rows(lesson)
    if not proc:
        proc = [["45'", "Vận dụng", f"Apply knowledge of {topic}", "Guide", "Present answers"]]
    _add_data_table(
        doc,
        ["Time", "Stage / Phần", "Content / Nội dung", "Teacher", "Students"],
        proc,
        theme=theme,
    )

    # ===== DETAILED SECTIONS =====
    _add_heading_custom(doc, "V. NỘI DUNG CHI TIẾT / DETAILED CONTENT", level=1, theme=theme)
    for i, sec in enumerate(lesson.get("sections") or [], 1):
        h_vi = sec.get("heading_vi") or ""
        h_en = sec.get("heading_en") or ""
        heading = h_vi or h_en or f"Phần {i}"
        if h_vi and h_en and h_vi != h_en:
            heading = f"{h_vi} / {h_en}"
        _add_heading_custom(doc, f"{i}. {heading}", level=2, theme=theme)

        if sec.get("duration"):
            _add_para(doc, f"Thời gian / Time: {sec.get('duration')}", size=11, italic=True, color=theme["orange"], space_after=4)

        c_vi = (sec.get("content_vi") or "").strip()
        c_en = (sec.get("content_en") or "").strip()
        if c_vi:
            _add_para(doc, "Nội dung (VI):", size=11, bold=True, color=theme["navy"], space_after=2)
            _add_para(doc, c_vi, size=13, color=GRAY, align="justify", first_line=True, space_after=6)
        if c_en:
            _add_para(doc, "Content (EN):", size=11, bold=True, color=theme["teal"], space_after=2)
            _add_para(doc, c_en, size=13, color=GRAY, align="justify", first_line=True, space_after=6)

        steps = sec.get("procedure_steps") or []
        if steps:
            _add_para(doc, "Các bước tiến hành:", size=11, bold=True, color=theme["navy"], space_after=2)
            for si, step in enumerate(steps, 1):
                _add_bullet(doc, f" {step}", bold_prefix=f"{si}.")
        if sec.get("teacher_activity"):
            _add_para(doc, f"Hoạt động GV (VI): {sec['teacher_activity']}", size=11, color=GRAY, space_after=2)
        if sec.get("teacher_activity_en"):
            _add_para(doc, f"Teacher (EN): {sec['teacher_activity_en']}", size=11, italic=True, color=GRAY, space_after=3)
        if sec.get("student_activity"):
            _add_para(doc, f"Hoạt động HS (VI): {sec['student_activity']}", size=11, color=GRAY, space_after=2)
        if sec.get("student_activity_en"):
            _add_para(doc, f"Students (EN): {sec['student_activity_en']}", size=11, italic=True, color=GRAY, space_after=3)
        if sec.get("expected_product"):
            _add_para(doc, f"Sản phẩm dự kiến: {sec['expected_product']}", size=11, italic=True, color=theme["teal"], space_after=3)
        questions = sec.get("questions") or []
        if questions:
            _add_para(doc, "Câu hỏi gợi mở:", size=11, bold=True, color=theme["navy"], space_after=2)
            for q in questions:
                _add_bullet(doc, f" {q}")
        _add_para(doc, "", space_after=4)

    # ===== VOCABULARY TABLE =====
    vocab = lesson.get("vocabulary") or []
    _add_heading_custom(doc, "VI. TỪ VỰNG / VOCABULARY", level=1, theme=theme)
    if vocab:
        rows = []
        for v in vocab:
            rows.append([
                v.get("term") or "",
                v.get("meaning_vi") or v.get("meaning") or "",
                v.get("example") or "",
            ])
        _add_data_table(doc, ["Term (EN)", "Meaning (VI)", "Example sentence"], rows, theme=theme)
    else:
        _add_para(doc, "Chưa có danh sách từ vựng.", size=11, italic=True, color=GRAY)

    # ===== ASSESSMENT =====
    _add_heading_custom(doc, "VII. KIỂM TRA – ĐÁNH GIÁ / ASSESSMENT", level=1, theme=theme)
    assessment = lesson.get("assessment") or {}
    if isinstance(assessment, str):
        _add_para(doc, assessment, size=12, align="justify", color=GRAY, first_line=True)
    else:
        methods = assessment.get("methods") or [
            "Observation during group work",
            "Oral questions / quick check",
            "Exit ticket or short quiz",
        ]
        criteria = assessment.get("criteria") or [
            "Participation and collaboration",
            "Correct use of key concepts / vocabulary",
            "Ability to apply knowledge to a new situation",
        ]
        _add_para(doc, "Phương pháp / Methods", size=12, bold=True, color=theme["teal"], space_after=4)
        for m in methods:
            _add_bullet(doc, f" {m}")
        _add_para(doc, "Tiêu chí / Criteria", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
        for c in criteria:
            _add_bullet(doc, f" {c}")

    # ===== HOMEWORK / REFLECTION =====
    _add_heading_custom(doc, "VIII. CỦNG CỐ – BÀI TẬP / CONSOLIDATION & HOMEWORK", level=1, theme=theme)
    homework = lesson.get("homework") or lesson.get("consolidation") or [
        f"Review key ideas about {topic}.",
        "Write 3–5 sentences in English using new vocabulary.",
        "Prepare one real-life example related to the topic for next lesson.",
    ]
    if isinstance(homework, str):
        homework = [homework]
    for h in homework:
        _add_bullet(doc, f" {h}")

    # ===== NOTES =====
    diff = (lesson.get("differentiation") or "").strip()
    if diff:
        _add_heading_custom(doc, "VIII.b. PHÂN HÓA – HỖ TRỢ / DIFFERENTIATION", level=1, theme=theme)
        _add_para(doc, diff, size=12, color=GRAY, align="justify", first_line=True, space_after=8)

    _add_heading_custom(doc, "IX. GHI CHÚ / NOTES", level=1, theme=theme)
    note = lesson.get("notes") or lesson.get("note") or (
        "Giáo viên điều chỉnh thời gian và hoạt động phù hợp năng lực lớp. "
        "Encourage bilingual interaction and experiential tasks."
    )
    _add_para(doc, note, size=11, italic=True, color=GRAY, align="justify")

    # Footer line
    _add_para(doc, "", space_after=8)
    _add_para(
        doc,
        "— EduBridge · Teacher Studio · Bilingual lesson plan —",
        size=9, italic=True, color=theme["orange"], align="center", space_before=12,
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _type_label(qtype: str) -> str:
    return {
        "multiple_choice": "Multiple choice",
        "true_false": "True / False",
        "fill_blank": "Fill in the blank",
        "essay": "Essay",
        "short_answer": "Essay",
    }.get(qtype or "", "Question")


def build_quiz_docx_bytes(quiz: dict) -> bytes:
    """Đề kiểm tra trang trí: header màu, khung thông tin HS, câu hỏi rõ ràng, đáp án cuối."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    theme = _theme_from(quiz)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    banner = doc.add_table(rows=1, cols=1)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = banner.cell(0, 0)
    _set_cell_shading(cell, _hex_of(theme["navy"]))
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EDUBRIDGE  ·  EXPERIENTIAL QUIZ (ENGLISH)")
    _set_run(r, size=16, bold=True, color=(255, 255, 255))
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(quiz.get("title") or quiz.get("topic") or "Quiz")
    _set_run(r2, size=13, bold=True, color=theme["orange"])

    _add_para(doc, "", space_after=4)
    _add_info_table(doc, [
        ("School", quiz.get("school") or "................................"),
        ("Grade", quiz.get("grade") or ""),
        ("Subject", quiz.get("subject") or ""),
        ("Topic", quiz.get("topic") or ""),
        ("Time", f"{quiz.get('time_limit_minutes') or 20} minutes"),
        ("Pass score", f"{quiz.get('passing_score') or 60}%"),
    ], theme=theme)

    if quiz.get("description"):
        _add_para(doc, quiz["description"], size=12, italic=True, color=GRAY, align="justify", space_after=8)

    info = doc.add_table(rows=2, cols=2)
    info.style = "Table Grid"
    _fill_cell(info.cell(0, 0), "Full name: ................................", size=12)
    _fill_cell(info.cell(0, 1), "Class: ....................", size=12)
    _fill_cell(info.cell(1, 0), "Date: ........................", size=12)
    _fill_cell(info.cell(1, 1), "Score: .......... / 10", size=12)
    doc.add_paragraph()

    questions = quiz.get("questions") or []
    for i, q in enumerate(questions, 1):
        qtype = q.get("question_type") or "multiple_choice"
        box = doc.add_table(rows=1, cols=1)
        box.style = "Table Grid"
        c = box.cell(0, 0)
        c.text = ""
        head = c.paragraphs[0]
        r_no = head.add_run(f"Question {i}  ")
        _set_run(r_no, size=13, bold=True, color=theme["navy"])
        r_tp = head.add_run(f"({_type_label(qtype)} · {int(q.get('points') or 1)} pt)")
        _set_run(r_tp, size=10, italic=True, color=theme["orange"])
        body = c.add_paragraph()
        r_q = body.add_run(q.get("question_text") or "")
        _set_run(r_q, size=13, color=GRAY)
        if qtype == "essay":
            for _ in range(4):
                line = c.add_paragraph()
                rr = line.add_run("................................................................................................")
                _set_run(rr, size=12, color=GRAY)
        elif qtype == "fill_blank":
            line = c.add_paragraph()
            rr = line.add_run("Answer: ................................")
            _set_run(rr, size=12, color=GRAY)
        elif qtype == "true_false":
            line = c.add_paragraph()
            rr = line.add_run("☐ True          ☐ False")
            _set_run(rr, size=12, color=theme["navy"])
        else:
            for opt in q.get("options") or []:
                op = c.add_paragraph()
                rr = op.add_run(f"☐  {opt}")
                _set_run(rr, size=12, color=GRAY)
        _set_cell_shading(c, "F8FBFF")
        doc.add_paragraph()

    doc.add_page_break()
    _add_heading_custom(doc, "ANSWER KEY — FOR TEACHERS / ĐÁP ÁN GV", level=1, theme=theme)
    ans_rows = []
    for i, q in enumerate(questions, 1):
        ans = q.get("correct_answer") or q.get("sample_answer") or "—"
        exp = q.get("explanation") or ""
        ans_rows.append([str(i), _type_label(q.get("question_type")), str(ans), str(exp)])
    if ans_rows:
        _add_data_table(doc, ["No.", "Type", "Answer / Key points", "Explanation"], ans_rows, theme=theme)
    _add_para(
        doc,
        "— EduBridge · Print page 1 for students; keep the answer key for teachers —",
        size=9, italic=True, color=theme["orange"], align="center", space_before=12,
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _worksheet_theme(data: dict):
    return _theme_from(data)


def _add_response_space(doc, task: dict, theme) -> None:
    rtype = (task.get("response_type") or "lines").lower()
    lines = int(task.get("response_lines") or 4)
    if rtype == "draw_box":
        _add_para(doc, "[ Draw here ]", size=11, italic=True, color=GRAY, space_after=4)
        for _ in range(6):
            _add_para(doc, "________________________________________________________", size=11, color=GRAY)
        return
    if rtype == "table" and isinstance(task.get("table_template"), dict):
        tbl = task["table_template"]
        headers = [str(h) for h in (tbl.get("headers") or ["", ""])]
        rows = tbl.get("rows") or [["", ""]]
        body = [[str(c) for c in row] for row in rows if isinstance(row, list)]
        _add_data_table(doc, headers, body, theme=theme)
        return
    if rtype == "mcq":
        for item in task.get("items") or []:
            _add_para(doc, str(item), size=12, color=GRAY, space_after=2)
            _add_para(doc, "☐ A    ☐ B    ☐ C    ☐ D", size=11, color=GRAY, space_after=6)
        return
    if rtype == "blanks":
        for item in task.get("items") or []:
            _add_para(doc, str(item).replace("___", "________________"), size=12, color=GRAY, space_after=6)
        return
    for item in task.get("items") or []:
        _add_para(doc, f"• {item}", size=12, color=GRAY, space_after=4)
    for _ in range(max(2, min(lines, 12))):
        _add_para(doc, "______________________________________________________________", size=11, color=GRAY)


def build_worksheet_student_docx_bytes(worksheet: dict) -> bytes:
    """Phiếu học tập — bản học sinh (không đáp án)."""
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    theme = _worksheet_theme(worksheet)
    st = worksheet.get("student") if isinstance(worksheet.get("student"), dict) else worksheet
    grade = worksheet.get("grade") or ""
    subject = worksheet.get("subject") or ""
    topic = worksheet.get("topic") or ""

    _add_para(doc, "EDUBRIDGE · STUDENT WORKSHEET", size=11, bold=True, color=theme["orange"], align="center", space_after=2)
    _add_heading_custom(doc, st.get("title") or f"Worksheet: {topic}", level=0, theme=theme)
    _add_info_table(doc, [
        ("Subject", subject),
        ("Grade", grade),
        ("Lesson / Topic", topic),
        ("Time", f"{worksheet.get('duration_minutes') or 15} minutes"),
    ], theme=theme)
    if st.get("header_fields"):
        _add_para(doc, "Name: _________________________    Class: _________________________    Date: ___________", size=11, color=GRAY, space_after=8)

    if st.get("learning_goal"):
        _add_para(doc, "Learning goal", size=12, bold=True, color=theme["teal"], space_after=4)
        _add_para(doc, st["learning_goal"], size=12, color=GRAY, align="justify", first_line=True, space_after=8)

    kws = st.get("key_words") or []
    if kws:
        rows = []
        for k in kws:
            if isinstance(k, dict):
                rows.append([k.get("term") or "", k.get("meaning_vi") or "", k.get("explanation_en") or ""])
        if rows:
            _add_para(doc, "Key words / Language support", size=12, bold=True, color=theme["teal"], space_after=4)
            _add_data_table(doc, ["English term", "Vietnamese", "Simple explanation"], rows, theme=theme)

    _add_heading_custom(doc, "Learning tasks", level=1, theme=theme)
    for task in st.get("tasks") or []:
        if not isinstance(task, dict):
            continue
        num = task.get("number") or ""
        title = task.get("title") or "Task"
        _add_heading_custom(doc, f"{num}. {title}".strip(), level=2, theme=theme)
        if task.get("instructions"):
            _add_para(doc, task["instructions"], size=12, bold=True, color=theme["navy"], space_after=4)
        if task.get("learning_material"):
            _add_para(doc, task["learning_material"], size=12, color=GRAY, align="justify", first_line=True, space_after=6)
        if task.get("time_minutes"):
            _add_para(doc, f"Time: about {task['time_minutes']} min", size=10, italic=True, color=theme["orange"], space_after=4)
        _add_response_space(doc, task, theme)
        _add_para(doc, "", space_after=6)

    refl = st.get("reflection") or []
    if refl:
        _add_heading_custom(doc, "Reflection", level=2, theme=theme)
        for r in refl:
            _add_bullet(doc, f" {r}")
            _add_para(doc, "______________________________________________", size=11, color=GRAY, space_after=4)

    if st.get("exit_question"):
        _add_heading_custom(doc, "Exit question", level=2, theme=theme)
        _add_para(doc, st["exit_question"], size=12, color=GRAY, space_after=4)
        _add_para(doc, "______________________________________________________________", size=11, color=GRAY)

    _add_para(doc, "— EduBridge · Student worksheet —", size=9, italic=True, color=theme["orange"], align="center", space_before=12)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_worksheet_teacher_docx_bytes(worksheet: dict) -> bytes:
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2)
    theme = _worksheet_theme(worksheet)
    tg = worksheet.get("teacher_guide") if isinstance(worksheet.get("teacher_guide"), dict) else {}
    st = worksheet.get("student") if isinstance(worksheet.get("student"), dict) else {}
    topic = worksheet.get("topic") or ""

    _add_para(doc, "EDUBRIDGE · TEACHER'S GUIDE & ANSWER KEY", size=11, bold=True, color=theme["orange"], align="center", space_after=4)
    _add_heading_custom(doc, st.get("title") or f"Worksheet: {topic}", level=0, theme=theme)

    for label, key in (
        ("Objectives", "objectives"),
        ("Common mistakes", "common_mistakes"),
        ("Suggested feedback", "suggested_feedback"),
        ("Assessment criteria", "assessment_criteria"),
    ):
        items = tg.get(key)
        if isinstance(items, list) and items:
            _add_para(doc, label, size=12, bold=True, color=theme["teal"], space_before=6, space_after=4)
            for it in items:
                _add_bullet(doc, f" {it}")

    _add_para(doc, "Organisation & timing", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
    _add_para(doc, f"Duration: {tg.get('suggested_duration') or worksheet.get('duration_minutes') or '—'}", size=11, color=GRAY)
    _add_para(doc, f"Mode: {tg.get('organisation') or '—'}", size=11, color=GRAY, space_after=6)

    if tg.get("teacher_says_en"):
        _add_para(doc, "Teacher says (English)", size=12, bold=True, color=theme["navy"], space_after=4)
        _add_para(doc, tg["teacher_says_en"], size=12, color=GRAY, align="justify", first_line=True, space_after=6)
    if tg.get("vietnamese_support"):
        _add_para(doc, "Hỗ trợ GV (Tiếng Việt)", size=12, bold=True, color=theme["navy"], space_after=4)
        _add_para(doc, tg["vietnamese_support"], size=12, color=GRAY, align="justify", first_line=True, space_after=8)

    keys = tg.get("answer_key") or []
    if keys:
        rows = []
        for k in keys:
            if isinstance(k, dict):
                rows.append([str(k.get("task") or ""), str(k.get("expected") or k.get("answer") or "")])
            else:
                rows.append(["", str(k)])
        _add_heading_custom(doc, "Answer key / Expected responses", level=1, theme=theme)
        _add_data_table(doc, ["Task", "Expected"], rows, theme=theme)

    if tg.get("transition"):
        _add_para(doc, "Transition", size=12, bold=True, color=theme["teal"], space_before=8, space_after=4)
        _add_para(doc, tg["transition"], size=12, color=GRAY, space_after=8)

    _add_para(doc, "— EduBridge · For teachers only —", size=9, italic=True, color=theme["orange"], align="center", space_before=12)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
